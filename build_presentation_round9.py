"""Editable, source-bound PowerPoint. Export with native PowerPoint using the companion script."""
from pathlib import Path
from hashlib import sha256
from datetime import datetime,timezone
import csv,json
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE,XL_LABEL_POSITION,XL_TICK_MARK,XL_LEGEND_POSITION,XL_TICK_LABEL_POSITION

ROOT=Path(__file__).resolve().parent;OUT=ROOT/'reports/presentation_round9';OUT.mkdir(exist_ok=True)
evidence=json.loads((ROOT/'reports/submission_round8/report_evidence.json').read_text(encoding='utf-8'))
metrics=evidence['metrics'];used={};extra={}
PERCENT_KEYS={'FIFO_CAP','Q_CAPTURE','GQ_CAPTURE','RANDOM_CAP','PROCESS_DEVIATION_CAP',
 'GQ_FPR','OOD05_FPR','GQ_INTERVENTION','OOD05_INTERVENTION','OOD05_CAPTURE'}
def metric(key):
    item=metrics[key];assert sha256((ROOT/item['source']).read_bytes()).hexdigest()==item['sha256']
    multiplier=100 if key in PERCENT_KEYS else 1
    used[key]={**item,'slide_multiplier':multiplier};return float(item['value'])*multiplier
def f(key,places=3):return f'{metric(key):,.{places}f}'
def source(path):
    extra[path]=sha256((ROOT/path).read_bytes()).hexdigest()
    return list(csv.DictReader((ROOT/path).open(encoding='utf-8-sig',newline='')))
groups=source('reports/oct03_round7/group_metrics.csv')
uncertainty=source('reports/oct03_round7/paired_uncertainty.csv')
q1=source('reports/oct03_pilot_r2/q1_comparison.csv')
source('reports/oct03_round7/queue_waits.csv') if (ROOT/'reports/oct03_round7/queue_waits.csv').exists() else None
for path in ['docs/sources/oct03_round4/employee_notice.txt','reports/oct03_evidence/final/official_contract.json',
 'docs/FUTURE_EVALUATION_CONTRACT.md','docs/DATA_REQUEST_ROUND8.md','future_evaluation_v2.py','tests/test_future_evaluation_v2.py']:
    extra[path]=sha256((ROOT/path).read_bytes()).hexdigest()

NAVY='14263D';BLUE='2864DC';GOLD='BE7F29';BG='F6F8FC';INK='172B43';MUTED='53657B';LINE='DCE3ED';WHITE='FFFFFF';LIGHT='EAF0FB'
prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
prs.core_properties.title='CastGuard — 제조 데이터 결합의 검사 가치';prs.core_properties.author='';prs.core_properties.last_modified_by=''
prs.core_properties.subject='검토용 발표 초안 · 탐색 근거와 현장 검증 조건';prs.core_properties.keywords=''
slides=[];geometry=[]
def color(value):return RGBColor.from_string(value)
def box(slide,x,y,w,h,fill=WHITE,line=None,radius=False):
    shape=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    shape.fill.solid();shape.fill.fore_color.rgb=color(fill)
    if line:shape.line.color.rgb=color(line);shape.line.width=Pt(.7)
    else:shape.line.fill.background()
    for effect in shape._element.xpath('.//a:effectRef'):effect.set('idx','0')
    return shape
def text(slide,x,y,w,h,value,size=20,fill=INK,bold=False,align=PP_ALIGN.LEFT):
    shape=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=shape.text_frame
    tf.word_wrap=True;tf.margin_left=tf.margin_right=0;tf.margin_top=tf.margin_bottom=0
    for i,line in enumerate(str(value).split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.alignment=align;p.space_after=Pt(7)
        for run in p.runs:run.font.name='맑은 고딕';run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=color(fill)
    geometry.append({'slide':len(prs.slides),'text':str(value),'x':x,'y':y,'w':w,'h':h,'size':size})
    return shape
def line(slide,x1,y1,x2,y2,c=LINE,width=1):
    obj=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));obj.line.color.rgb=color(c);obj.line.width=Pt(width)
def new(title,kicker,subtitle='',source_ids='',note='',seconds=0,appendix=False):
    slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=color(BG)
    box(slide,0,0,.12,7.5,BLUE)
    text(slide,.58,.32,11.8,.28,('APPENDIX  /  ' if appendix else 'CASTGUARD  /  ')+kicker,11,BLUE,True)
    text(slide,.58,.82,12.0,.7,title,29,INK,True)
    if subtitle:text(slide,.6,1.55,12.0,.58,subtitle,15,MUTED)
    line(slide,.58,6.9,12.75,6.9)
    text(slide,.6,7.02,11.45,.25,source_ids+'  |  탐색자료 · 현장 효과 미입증',9,MUTED)
    text(slide,12.05,7.01,.65,.25,f'{len(prs.slides):02d}',10,MUTED,align=PP_ALIGN.RIGHT)
    slide.notes_slide.notes_text_frame.text=note+'\n\n근거: '+source_ids+'\n이 슬라이드의 수치는 과거 탐색 결과이며 새 독립 검증이 아닙니다.'
    slides.append({'number':len(prs.slides),'title':title,'section':kicker,'appendix':appendix,'practice_seconds':seconds,'speaker':note,'source_ids':source_ids})
    return slide
def card(slide,x,y,w,h,title,body,accent=BLUE):
    box(slide,x,y,w,h,WHITE);box(slide,x,y,.055,h,accent)
    text(slide,x+.2,y+.2,w-.4,.6,title,20,accent,True)
    text(slide,x+.2,y+1.0,w-.4,h-1.1,body,18)
def table(slide,x,y,width,headers,rows,widths=None,row_h=.52,font=16):
    widths=widths or [1/len(headers)]*len(headers)
    for i,row in enumerate([headers]+rows):
        xx=x;hh=row_h
        for j,(value,fraction) in enumerate(zip(row,widths)):
            ww=width*fraction;box(slide,xx,y+i*hh,ww-.02,hh-.025,NAVY if i==0 else WHITE)
            text(slide,xx+.13,y+i*hh+.105,ww-.25,hh-.12,str(value),font,WHITE if i==0 else INK,i==0)
            xx+=ww
def bar_chart(slide,x,y,w,h,labels,values,minimum=0,maximum=25,number_format='0.0',series_name='포착률 (%)'):
    data=CategoryChartData();data.categories=labels;data.add_series(series_name,values)
    chart=slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED,Inches(x),Inches(y),Inches(w),Inches(h),data).chart
    chart.has_legend=False;chart.has_title=False
    chart.chart_style=10
    va=chart.value_axis;va.minimum_scale=minimum;va.maximum_scale=maximum;va.major_unit=5
    va.tick_labels.font.name='맑은 고딕';va.tick_labels.font.size=Pt(12);va.tick_labels.number_format='0'
    va.tick_label_position=XL_TICK_LABEL_POSITION.LOW
    va.major_tick_mark=XL_TICK_MARK.NONE;va.has_major_gridlines=True;va.major_gridlines.format.line.color.rgb=color(LINE)
    ca=chart.category_axis;ca.reverse_order=True;ca.tick_labels.font.name='맑은 고딕';ca.tick_labels.font.size=Pt(15)
    ca.tick_label_position=XL_TICK_LABEL_POSITION.LOW
    ca.major_tick_mark=XL_TICK_MARK.NONE
    plot=chart.plots[0];plot.gap_width=75;plot.has_data_labels=True
    plot.data_labels.position=XL_LABEL_POSITION.OUTSIDE_END;plot.data_labels.number_format=number_format
    plot.data_labels.font.name='맑은 고딕';plot.data_labels.font.size=Pt(15)
    chart.series[0].invert_if_negative=False
    for i,point in enumerate(chart.series[0].points):
        point.format.fill.solid();point.format.fill.fore_color.rgb=color(GOLD if values[i]<0 else BLUE)
        point.format.line.fill.background()
    return chart

s=new('제조 데이터 결합의 검사 가치를 검증하다','제한된 검사 용량 · 다종 데이터 · 재현 가능한 판단',
 'CastGuard  |  제6회 K-인공지능 제조데이터 분석 경진대회 검토용 발표 초안',
 'S01 공식 과제·양식 / S02 현재 보고서',
 '저희는 제한된 검사 자원에서 어떤 Shot을 먼저 확인할지에 초점을 맞췄습니다. 품질 위험도와 설비 상태를 결합하고 검사 대기까지 계산했습니다. 현재 데이터에서는 복잡한 결합의 일관된 추가가치를 입증하지 못했습니다. 오늘은 그 판단의 근거와, 현장 검증으로 연결하기 위해 실제로 준비한 체계를 설명하겠습니다.',30)
text(s,.68,2.55,10.9,1.3,'품질 위험도와 설비 상태를\n같은 검사 자원 위에서 비교했습니다.',34,INK,True)
box(s,.7,4.6,11.9,1.45,NAVY)
text(s,.98,4.84,11.3,.43,'현재 결론  ·  결합의 현장 우위는 미입증',25,WHITE,True)
text(s,.98,5.4,11.3,.4,'고정 비교 → 실패 원인 → 입력·검사 제약 → 독립 실증 계획',18,'C4D3E9')

s=new('검사 자원이 의사결정을 제한한다','01 기업문제',
 '실제 사업장 비용·검사시간은 미확인입니다. 아래는 평가할 의사결정의 정의입니다.',
 'S02 §1·§4 / S03 queue protocol',
 '한 번에 모든 제품을 검사할 수 없다고 가정하면 위험도 점수만 높여서는 충분하지 않습니다. 누락된 입력은 어떻게 보류할지, 설비 경보가 품질 검사를 얼마나 밀어내는지, 대기 중 검사가 유효기간을 넘는지도 같이 봐야 합니다. 저희의 목적은 같은 검사 자원으로 포착하는 불량을 비교하면서 정상 경보와 미처리 부담을 함께 제시하는 것입니다.',45)
card(s,.68,2.35,3.8,3.7,'생산 흐름','Shot 도착\n공정 입력의 도착시점\n설비 상태·품질 위험도')
card(s,4.77,2.35,3.8,3.7,'공통 검사 제약','20% 누적 용량 상한\n입력 누락 보류 우선\n실제 검사량·잔량 계산')
card(s,8.86,2.35,3.8,3.7,'함께 볼 결과','관측 불량 포착률\n정상 경보·보류 부담\n대기·미처리·미관측',GOLD)

s=new('두 데이터의 역할을 연결하고, 세 번째의 적용 범위를 구분했다','02 데이터 통합',
 '물리적 설비 계보·단위·실제 수신시각은 원본 가이드와 현장 로그 확인이 필요합니다.',
 'S04 lineage audit / S02 §1·§2',
 '주 데이터 #42는 주조 품질, 보조 #41은 설비 상태와 이력입니다. 파생한 구간·Shot 키로 대응되는 행을 확인했지만, 공정값이 같은 부분은 새로운 독립 센서 정보가 아닙니다. #40은 별도 데이터 내부에서 변수 가치만 확인했으며 #42로 직접 붙이거나 성능 전이 근거로 쓰지 않았습니다. 파일 번호는 프로젝트 식별자이고 포털 식별자는 별도로 기록했습니다.',50)
card(s,.68,2.28,3.8,3.95,'주 데이터  #42',f'주조 품질·공정값\n중복 정리 후 {f("Q_N",0)}행\n관측 불량 {f("Q_POS",0)}행\n품질 위험도 학습')
card(s,4.77,2.28,3.8,3.95,'보조 데이터  #41',f'설비 상태·운전 이력\n대응 키 {f("Q_N",0)}개\n공정값 {f("JOIN_VALUES",0)}개 일치\n상태 경보·이력 결합')
card(s,8.86,2.28,3.8,3.95,'별도 보완  #40','금형온도 입력 제거실험\n같은 #40 내부의 변수 가치\n#42와 행 단위 결합 없음\n외부 전이 성능 미입증',GOLD)

s=new('같은 흐름과 같은 용량에서 정책을 비교했다','03 검증 설계',
 '개발 선택은 안쪽 시간분할에서 수행했습니다. 기존 자료는 이미 노출되어 탐색 평가입니다.',
 'S03 queue protocol / S05 group_metrics / S06 exposure',
 '시간순 안쪽 분할로 후보를 고정하고 바깥쪽 네 구간의 생산 흐름을 순서대로 재생했습니다. 모든 정책에 같은 누적 검사 상한과 보류 규칙을 적용했습니다. 상한이 같아도 유효 대기열의 차이로 실제 검사 수는675건과677건으로 달랐습니다. 품질이 관측된 2879행의 양성505개가 포착률 분모이고, 품질 미관측516행도 검사 부담에는 포함했습니다.',50)
table(s,.7,2.35,7.4,['비교 축','고정한 조건'],[
 ['후보 선택','안쪽 시간순 개발 분할'],['비교 정책','FIFO · 무작위 · 공정이탈 · Q · GQ'],
 ['자원','누적20% 상한 · 동일 hold 우선순위'],['실제 검사','Q/FIFO 675건 · GQ 677건'],['평가 한계','이미 본4구간 · 현장 타이밍 미확인']],widths=[.25,.75],row_h=.67,font=17)
text(s,8.55,2.48,3.8,.7,f('QUEUE_N',0),39,BLUE,True);text(s,8.55,3.2,3.7,.4,'전체 생산 기록',17,MUTED)
text(s,8.55,4.06,3.8,.7,f('QUEUE_POS',0),39,BLUE,True);text(s,8.55,4.78,3.7,.8,'품질 관측 양성 분모\n품질 미관측 516행 별도',17,MUTED)

s=new('단순 기준 대비 모델의 추가가치는 확인되지 않았다','04 동일 용량 결과',
 '관측 불량 포착률 (%) = 검사된 관측 양성 / 관측 양성505개',
 'S05 group_metrics / S07 baseline results',
 'FIFO의22.376%가 Q20.871%, GQ20.238%보다 높았습니다. 과거 제품빈도 기준은 구간마다 제품이 하나라 점수가 상수였고 FIFO와 검사 행·시점이 완전히 같았습니다. 두 개의 독립적인 성공 기준으로 세지 않았습니다. 무작위는100회 평균이며 이 표만으로 FIFO의 현장 우위를 주장하지 않습니다. 다음 장에서 결합 전후와 구간별 원인을 구분해 설명하겠습니다.',55)
bar_chart(s,.7,2.22,8.0,4.2,['FIFO = 제품빈도','Q 품질 단독','GQ 상태+품질','무작위100회 평균','공정이탈'],
 [metric('FIFO_CAP'),metric('Q_CAPTURE'),metric('GQ_CAPTURE'),metric('RANDOM_CAP'),metric('PROCESS_DEVIATION_CAP')],number_format='0.000')
card(s,9.02,2.47,3.56,3.5,'해석의 경계','같은 과거4구간\nQ/GQ는5 seed 평균\n무작위 범위는 조건부\n현장 우위·동등성 미입증',GOLD)

s=new('결합 전후의 성과를 같은 평가 안에서 확인했다','05 다종 결합의 추가가치',
 '첫 행은 품질 예측 AP, 두 번째는 순차 검사 포착률입니다. 서로 다른 평가를 합산하지 않습니다.',
 'S02 §2·§3 / S08 ablation_summary / S05 group_metrics',
 '융합 요구에 대응하기 위해 데이터 개수만 나열하지 않고 결합 전후를 제시합니다. 같은 within-run 평가에서 주 품질모델A의AP는.2614, 이력 결합B는.2668로 약.00535 올라 사전 기준.02에 미달했습니다. 별도의 같은 순차 검사 평가에서는 Q에서GQ로 결합했을 때 관측 불량 포착률이.634퍼센트포인트 낮았습니다. 두 결과는 평가가 다르므로 합산하지 않으며, 통합 성공이라는 표현을 사용하지 않습니다.',50)
table(s,.7,2.45,11.9,['평가·전후','주 데이터 / 단독','보조 결합 후','판정'],[
 ['within_run AP',f'A  {f("A_AP",4)}',f'B  {f("B_AP",4)}','증가 기준 +0.02 미달'],
 ['순차 검사 포착률',f'Q  {f("Q_CAPTURE")}% ',f'GQ  {f("GQ_CAPTURE")}%','추가가치 채택 실패']],
 widths=[.23,.23,.23,.31],row_h=.9,font=18)
box(s,.73,5.56,11.84,.66,LIGHT);text(s,.95,5.72,11.4,.36,'통합 구현과 통합 성능 개선을 구분했습니다. 원본·분할·실패 결과를 보존합니다.',18,BLUE,True)

s=new('오경보 감소와 검사 부담을 함께 봐야 한다','06 운영 부담',
 'OOD05는 학습 분포 이탈을 보류하는 비교안입니다. 운영 정책으로 채택하지 않았습니다.',
 'S03 outer_policy_metrics / S02 §4',
 '분포 이탈 보류를 추가하면 정상 자동 경보율은 낮아집니다. 그러나 경보 또는 보류에 해당하는 정상 기록이 늘고 미처리 보류는241개에서988개가 됐습니다. 품질 포착률도 개선되지 않았습니다. 따라서 오경보율 하나만 보여 안전한 정책으로 소개하지 않고, 검사 자원을 함께 사용하는 전체 부담을 공개했습니다.',45)
table(s,.7,2.32,11.9,['같은4구간·20% 상한','GQ','GQ + OOD05 보류'],[
 ['정상 자동 FPR',f'{f("GQ_FPR")}% ',f'{f("OOD05_FPR")}%'],
 ['정상 경보 ∪ 보류',f'{f("GQ_INTERVENTION")}% ',f'{f("OOD05_INTERVENTION")}%'],
 ['미처리 보류',f'{f("GQ_BACKLOG",0)}개',f'{f("OOD05_BACKLOG",0)}개'],
 ['관측 불량 포착률',f'{f("GQ_CAPTURE")}% ',f'{f("OOD05_CAPTURE")}%']],widths=[.46,.25,.29],row_h=.69,font=19)
text(s,.78,6.1,11.7,.44,'상태 경보 감소 ≠ 품질 개선 또는 실제 비용 절감',22,GOLD,True)

s=new('Q의 전체 손실은 한 구간에 집중됐다','07 오류 원인',
 'FIFO 대비 포착 양성 수 차이 (건,5 seed 평균) · 구간별 양성 분모 158 / 5 / 44 / 298',
 'S05 group_metrics / S09 paired diagnosis',
 'Q와FIFO의 전체 차이는 평균 마이너스7.6건입니다. 이를 구간별로 나누면 run1에서14건 손실이고 나머지에서6.4건 이득입니다. run1을 제외하면 방향이 바뀝니다. run1은 앞쪽에 양성이 더 많았고 FIFO가 오래 기다린 초기 행을 검사하는 구조였습니다. 이는 저장된 행동의 산술적 설명이며 공정의 인과 원인을 밝힌 것은 아닙니다.',55)
run_rows=[next(r for r in groups if r['dimension']=='run_id' and r['group']==str(i) and r['policy']=='Q') for i in range(1,5)]
bar_chart(s,.7,2.28,7.85,4.12,[f'run {i}' for i in range(1,5)],[float(r['delta_hits_vs_FIFO']) for r in run_rows],minimum=-20,maximum=10,number_format='+0.0;-0.0;0',series_name='Q−FIFO 포착 양성 (평균 건)')
card(s,8.95,2.54,3.64,3.5,'전체 −7.6건','run1  −14.0건\n나머지  +6.4건\nrun1 제외 시 부호 반전\n구간 평균 일반화 금지',GOLD)

s=new('불확실성과 대기 한계가 현장 우위 판단을 막는다','08 결론의 강도',
 '4개 run의 조건부 재표본 범위입니다. 독립 검증의95% 일반화 신뢰구간이 아닙니다.',
 'S10 paired_uncertainty / S09 waits & exposure',
 '고정된 과거 행동을 네 구간 단위로 재표본하면 Q와GQ의 FIFO 대비 차이 범위가 모두0을 포함합니다. 유의한 우위도 동등성도 주장할 수 없습니다. 또한FIFO의 품질검사 평균 대기는310기록으로Q의142기록보다 길었습니다. 실제 시각이 없으므로 분 단위로 환산할 수 없습니다. 평가에 적격한 품질4613행은 모두 과거test에 노출됐습니다.',50)
text(s,.78,2.24,7.5,.4,'FIFO 대비 포착률 차이 (%p)',18,INK,True)
x0,x1,lo,hi=.99,8.07,-10,5
scale=lambda v:x0+(v-lo)/(hi-lo)*(x1-x0)
for tick in [-10,-5,0,5]:
    xx=scale(tick);line(s,xx,2.96,xx,5.48,NAVY if tick==0 else LINE,1.5 if tick==0 else .6)
    text(s,xx-.3,5.58,.6,.35,str(tick),13,MUTED,align=PP_ALIGN.CENTER)
for i,p in enumerate(['Q','GQ']):
    row=next(r for r in uncertainty if r['policy']==p and r['method']=='whole_run_cluster_exact_enumeration')
    point,lower,upper=[float(row[k]) for k in ['point_pp','lower_pp','upper_pp']];yy=3.45+i*1.3
    text(s,.77,yy-.65,7.4,.42,f'{p}: {point:+.3f}  [{lower:+.3f}, {upper:+.3f}]',18,INK,True)
    line(s,scale(lower),yy,scale(upper),yy,BLUE,3)
    for edge in [lower,upper]:line(s,scale(edge),yy-.1,scale(edge),yy+.1,BLUE,1.5)
    box(s,scale(point)-.055,yy-.075,.11,.15,BLUE)
card(s,8.95,2.45,3.64,3.72,'함께 남는 한계',f'품질검사 평균 대기\nFIFO {f("R7_WAIT_FIFO",0)} / Q {f("R7_WAIT_Q",0)}기록\n실제 시간·만료 미확인\n적격 {f("R7_EXPOSED_TEST",0)}행 과거노출',GOLD)

s=new('실제 자료가 오면 평가 전에 잘못된 입력부터 걸러낸다','09 구현한 기능',
 '오프라인 계약 검증기입니다. 구조 통과가 원천 진실성·현장 성능 인증을 뜻하지 않습니다.',
 'S11 future evaluation contract / S12 v2 regression tests',
 '평가를 재개할 때 같은 문제가 반복되지 않도록 입력계약을 코드로 구현했습니다. 입력과 라벨이 언제 이용 가능했는지, 정책과 독립적인 감사표본인지, 같은 검사 자원과 완료시점·만료를 적용했는지를 검사합니다. 이번 코드 검토에서는 선행0이나NA라는 식별자가CSV해석으로 변형되는 문제와 빈 구간ID도 고쳤습니다.109개 테스트는 기능 검증이며 성능이 올랐다는 근거는 아닙니다.',55)
card(s,.7,2.35,3.78,3.78,'① 정보시점','생산·센서 도착\n점수 결정·검사 완료\n라벨 이용 가능 시점\n미래 입력·정답 차단')
card(s,4.79,2.35,3.78,3.78,'② 독립 표본','전체 생산행 유지\n고정된 감사 추출\n모든 선택 표본의 정답\n정책 선택 표본 혼입 차단')
card(s,8.88,2.35,3.78,3.78,'③ 공통 제약','공통 슬롯·자원\n보류·대기공간·만료\n완료 가능한 검사만 인정\n사람의 원천 검토 필요',GOLD)

s=new('검사·대기·미관측을 함께 평가하는 실험 체계를 만들었다','10 기여와 차별점',
 '새 알고리즘의 세계 최초성이나 외부 전이 성능을 주장하지 않습니다.',
 'S02 current report / S09 diagnosis / S11 contract',
 '현재 제안의 기여는 제조 데이터 결합을 실제 의사결정 조건에서 반증 가능하게 만드는 실험 체계입니다. 품질과 상태를 연결한 전후 비교, 검사 자원을 공유할 때의 기회비용, 정답이 없는 구간을 숨기지 않는 분석을 구현했습니다. 실제 기업문제 해결이나 경제효과는 입증하지 못했으며 경쟁력의 핵심 공백으로 남아 있습니다.',40)
table(s,.7,2.3,11.9,['구현한 기여','검토 가능한 근거','아직 입증하지 못한 것'],[
 ['다종 결합 전후 비교','A→B · Q→GQ · OOD 비교','일관된 품질/운영 추가가치'],
 ['검사 자원 전체 회계','경보·보류·지연·미처리','실제 현장 비용·ROI'],
 ['실패 원인의 분해','구간·동점·대기·평가노출','물리적 원인·다른 공장 전이'],
 ['실증의 입력 계약','시각·표본·용량·만료 검증','원천 진실성·실제 현장 효용']],widths=[.29,.34,.37],row_h=.77,font=17)

s=new('다음 실증은 같은 생산흐름에 고정한 정책을 비교한다','11 미래 평가 계획',
 '아직 수행하지 않은 계획입니다. 실제 기준값·표본 규모·운영조건을 정하고 새 라벨을 보기 전에 고정합니다.',
 'S13 future statistical analysis plan / S11 data request',
 '새 생산 전체 모집단에서 입력과 정책을 먼저 고정하고 정책과 독립적인 감사표본을 뽑습니다. 같은 시점의 FIFO·품질 단독·결합 정책을 동일 자원으로 재생합니다. 감사가 전수가 아니면 추출확률을 사용해 포착률을 추정하고 같은 표본에서 정책 차이를 비교합니다. 필요한 표본 수나 성공 기준은 사업상 최소 가치와 새 수집 설계를 확인해 라벨 열람 전에 확정해야 합니다.',45)
table(s,.7,2.33,11.9,['단계','고정하거나 확보할 것'],[
 ['1. 새 모집단','기존 노출자료와 분리된 전체 생산 흐름·원본 키'],
 ['2. 실제 제약','수신/라벨 시각 · 검사 자원 · 처리시간 · 대기/만료'],
 ['3. 독립 관측','전수 또는 사전 고정 확률 감사 · 별도 자원'],
 ['4. 쌍을 이룬 비교','동일 감사표본 · 동일 자원 · 고정 정책의 차이'],
 ['5. 판정','사전 효용 기준·부담 상한 · 구간 불확실성 · 실패 공개']],widths=[.23,.77],row_h=.65,font=17)

s=new('측정한 실패를 다음 실증의 구체적 조건으로 바꿨다','12 제안의 현재 판단',
 '현재 데이터에서 운영 우위는 미입증입니다. 검증된 기능과 현장 효과를 분리해 제안합니다.',
 'S02 current report / S13 analysis plan',
 '현재 데이터에서는 품질모델과 설비 상태 결합이 단순 기준보다 일관되게 낫다는 증거가 없습니다. 자동정지나 비용 절감을 약속할 수는 없습니다. 같은 자원 위에서 결합의 효과를 검증하는 코드, 실패를 설명하는 근거, 새 실증에 필요한 입력계약을 제공합니다. 다음 판단은 새 독립 품질 관측과 실제 대기 제약에서 내리겠습니다.',30)
card(s,.7,2.42,3.78,3.52,'현재 유지','재현 가능한 비교 코드\n분모·누락·실패의 공개\n검사 의사결정 중심의 문제')
card(s,4.79,2.42,3.78,3.52,'현재 보류','모델/FIFO의 현장 우위\n자동 생산중지 적용\n통합 성공·실제 ROI',GOLD)
card(s,8.88,2.42,3.78,3.52,'다음 근거','새 생산·독립 품질 관측\n실제 시각·검사 제약\n라벨 열람 전 고정 평가')

s=new('특정 불량 전문화도 미래 개발 구간에서 실패했다','A Q1 상세',
 'Short_Shot 포착률 · 저장된5 seed 평균 · 각 예산은 구간 전체 순위검사이며 순차 queue와 다릅니다.',
 'S14 q1_comparison / selection',
 'Q1은20% 개발 기준으로 후보를 고정했습니다.10·30%는 저장된 보조 비교이고 이 표를 보고 새 후보를 선택하지 않았습니다. 두 미래 validation 구간 모두 세 예산에서 기존 A보다 낮았습니다. 개발 실패에 따라 Q1 test는 실행하지 않았습니다. 다른 희소 불량 유형은 미실행으로 남겼습니다.',appendix=True)
rows=[]
for fold in ['2','6']:
    for budget in [.1,.2,.3]:
        group=[r for r in q1 if r['scheme']=='forward_run' and r['fold']==fold and float(r['budget'])==budget]
        assert len(group)==5
        avg=lambda k:sum(float(r[k]) for r in group)/len(group)
        rows.append([fold,f'{budget:.0%}',str(int(float(group[0]['positive']))),f'{avg("capture_base")*100:.3f}%',f'{avg("capture")*100:.3f}%',f'{avg("gain")*100:+.3f}%p'])
table(s,.7,2.3,11.9,['fold','검사 예산','양성 분모','기존 A','Q1 전문화','Q1−A'],rows,widths=[.1,.16,.16,.18,.18,.22],row_h=.55,font=17)

s=new('G1은 오경보를 줄였지만 목표와 최악 구간을 통과하지 못했다','B Gate 상세',
 '상태 탐지와 품질 포착은 다른 결과입니다. 관측/전체 예열 분모와 정상 FPR을 함께 보고합니다.',
 'S15 pilot RESULTS / G1 metrics',
 'run_holdout에서는 정상 FPR이14.021%에서7.830%로 줄었지만2% 목표를 넘었고 최악 run은84.081%였습니다. 관측 예열은20개 중17~18개, 전체 예열은21개입니다. forward에서는 FPR.458%로 낮았지만 단순 규칙이.170%에서 같은11개 에피소드를 찾았습니다. 예열 구간의 품질 정답이 없으므로 상태 탐지를 불량 감소로 바꾸지 않습니다.',appendix=True)
table(s,.7,2.35,11.9,['평가','기존 → 보수 정책 FPR','관측 탐지 / 전체','남은 문제'],[
 ['run_holdout','14.021% → 7.830%','17~18 /20 (전체21)','정상 목표2% 미달'],
 ['forward','3.080% → 0.458%','11 /15 (전체16)','단순 규칙0.170%,11/15']],widths=[.2,.29,.26,.25],row_h=.91,font=17)
text(s,.8,5.48,11.8,.8,'정상 분모: run_holdout 4,613 Shot / forward 2,357 Shot\n최악 run FPR: 84.081% / 실제 생산정지·불량 감소는 미측정',20,GOLD,True)

s=new('분모와 단위를 바꾸면 다른 결론이 된다','C 지표 정의',
 'Q = 품질 단독 / GQ = 상태 Gate + 품질 / OOD05 = 분포 이탈 보류안',
 'S03 protocol / S05 group_metrics / S09 diagnosis',
 '품질 포착률의 분모는 관측 양성505개이며 전체 생산3395행에 정답이 모두 있는 것은 아닙니다. 정상 자동 FPR의 분모는 상태 정상3191행입니다. 대기는 기록 수 단위이고5 seed 평균은 표본 수를 다섯 배로 늘리지 않습니다. 제품빈도와 FIFO는 이번 구간에서 같은 행동이므로 중복 성공으로 세지 않습니다.',appendix=True)
table(s,.7,2.25,11.9,['지표','정의·분모','해석 주의'],[
 ['품질 포착률','검사된 관측 양성 /505','전체 생산 불량 포착률 아님'],
 ['정상 자동 FPR','상태정상 경보 /3,191','정상 경보∪보류와 구분'],
 ['미관측 품질','516 /3,395 생산 기록','정상 또는 불량0으로 대체 금지'],
 ['평균 대기','검사된 품질행의 기록수 차이','분·시간 환산 불가'],
 ['5 seed 평균','고정 표본에 대한 반복 학습','독립 표본·일반화CI 아님'],
 ['검사 상한20%','시점별 누적 허용 슬롯','실제 이용량675~677건']],widths=[.21,.4,.39],row_h=.56,font=16)

s=new('공식 여섯 항목에 근거와 공백을 연결했다','D 심사 요구 대응',
 '과제 배점과 양식의2·3장 배점이 다릅니다. 우선순위 미확인 상태를 그대로 표시합니다.',
 'S01 employee notice & template / S16 competition review',
 '기업문제와 통합설계, 모델과 통합성능, 영향요인과 오류, 현장 적용, 창의성과 확장, 재현성의 여섯 항목을 모두 다룹니다. 통합성능과 실제 기업 효과는 가장 큰 공백입니다. 좋은 발표나 많은 테스트가 이를 대신하지는 못합니다. 최종 포털 규격과 블라인드 검수는 별도로 완료해야 합니다.',appendix=True)
table(s,.7,2.26,11.9,['항목','과제 / 양식','연결한 근거','남은 공백'],[
 ['문제·통합','20 /20','본문2~3','물리적 계보·실제 사업장'],['모델·성능','35 /40','본문4~7, 부록A/B','통합 추가가치 미입증'],
 ['영향·오류','15 /10','본문8~9','인과 원인·실측 상호작용'],['현장 적용','10 /10','본문7,10,12','실제 제약·ROI'],
 ['창의·확장','10 /10','본문11~12','독창성·외부 전이 성능'],['코드·재현','10 /10','원본·모델·계약·시험','새 환경 설치·전체 재학습']],widths=[.2,.18,.31,.31],row_h=.56,font=16)

assert len(slides)==17
path=OUT/'CastGuard_presentation_draft.pptx';prs.save(path)
sources={
 'S01':'docs/sources/oct03_round4/employee_notice.txt; reports/oct03_evidence/final/official_contract.json',
 'S02':'reports/submission_round8/CastGuard_current_review.md',
 'S03':'docs/OCT03_QUEUE_PROTOCOL.md; reports/oct03_queue/outer_policy_metrics.csv',
 'S04':'reports/oct03_evidence/final/lineage_audit.json',
 'S05':'reports/oct03_round7/group_metrics.csv','S06':'reports/oct03_round7/evaluation_exposure.json',
 'S07':'reports/submission_round6/baselines_final','S08':'reports/oct02/ablation_summary.csv',
 'S09':'reports/oct03_round7/DIAGNOSTIC.md','S10':'reports/oct03_round7/paired_uncertainty.csv',
 'S11':'docs/FUTURE_EVALUATION_CONTRACT.md; docs/DATA_REQUEST_ROUND8.md',
 'S12':'future_evaluation_v2.py; tests/test_future_evaluation_v2.py',
 'S13':'docs/FUTURE_STATISTICAL_PLAN_ROUND9.md',
 'S14':'reports/oct03_pilot_r2/q1_comparison.csv; reports/oct03_pilot_r2/selection.json',
 'S15':'reports/oct03_pilot_r2/RESULTS.md','S16':'reports/presentation_round9/COMPETITION_REVIEW.md'}
for sourcestring in sources.values():
    for item in sourcestring.split('; '):
        target=ROOT/item
        if target.is_file():extra[item]=sha256(target.read_bytes()).hexdigest()
manifest={'created_at':datetime.now(timezone.utc).isoformat(),'slides':slides,'numeric_metrics':used,'source_hashes':extra,
 'source_index':sources,'core_slides':13,'appendix_slides':4,'practice_seconds':sum(s['practice_seconds'] for s in slides),
 'practice_timing_is_not_official':True,'new_performance_evidence':False,'pptx_sha256':sha256(path.read_bytes()).hexdigest(),
 'editable_text_charts_and_notes':True,'presenter_identity_blank':True,'renderer':'Microsoft PowerPoint native export required'}
(OUT/'slide_evidence.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
(OUT/'layout_geometry.json').write_text(json.dumps(geometry,ensure_ascii=False,indent=2),encoding='utf-8')
script=['# CastGuard 발표자 원고 — 회차9','',
 '검토용 본문13장+부록4장. 본문 연습 배분은 총'+str(manifest['practice_seconds'])+'초이며 공식 발표시간을 확인한 값이 아니다. 실제 제한에 따라 분량을 조정한다. 이름·소속·서명은 만들지 않았다. 슬라이드 노트에도 같은 원고가 있다.','']
for item in slides:
    script += [f'## {item["number"]:02d}. {item["title"]}',f'연습 배분: {item["practice_seconds"]}초' if not item['appendix'] else '질문 시 사용하는 부록',item['speaker'],'',f'근거: {item["source_ids"]}','']
script+=['## 근거 경로','']+[f'- {k}: `{v}`' for k,v in sources.items()]
(OUT/'SPEAKER_NOTES.md').write_text('\n'.join(script),encoding='utf-8')
(ROOT/'requirements-presentation.txt').write_text('python-pptx==1.0.2\npypdf==6.10.0\n# Native PDF/PNG export: Microsoft PowerPoint on Windows (installed application).\n',encoding='utf-8')
print(json.dumps({'slides':len(slides),'numeric_bindings':len(used),'sources':len(extra),'pptx_bytes':path.stat().st_size,'practice_seconds':manifest['practice_seconds']},ensure_ascii=False))
