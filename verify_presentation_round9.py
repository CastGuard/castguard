"""Verify native slide/PDF contents and source units without rebuilding authored slides."""
from pathlib import Path
from historical_sources import resolve_historical_source
from hashlib import sha256
from datetime import datetime,timezone
import csv,json,re,zipfile,math
from pptx import Presentation
from pypdf import PdfReader
ROOT=Path(__file__).resolve().parent;OUT=ROOT/'reports/presentation_round9'
def digest(p):return sha256(p.read_bytes()).hexdigest()
def verify():
    meta=json.loads((OUT/'slide_evidence.json').read_text(encoding='utf-8'))
    resolved_sources={path:resolve_historical_source(ROOT,path,h).relative_to(ROOT).as_posix() for path,h in meta['source_hashes'].items()}
    pptx=OUT/'CastGuard_presentation_draft.pptx';pdf=OUT/'CastGuard_presentation_draft.pdf'
    assert digest(pptx)==meta['pptx_sha256']
    presentation=Presentation(pptx);pages=PdfReader(pdf)
    assert len(presentation.slides)==len(pages.pages)==17
    texts=['\n'.join(shape.text for shape in slide.shapes if shape.has_text_frame) for slide in presentation.slides]
    pdftext=[page.extract_text() for page in pages.pages]
    for i,slide in enumerate(presentation.slides):
        assert slide.notes_slide.notes_text_frame.text.startswith(meta['slides'][i]['speaker'])
        assert meta['slides'][i]['title'] in texts[i]
        squash=lambda v:re.sub(r'\s+','',v)
        assert squash(meta['slides'][i]['title']) in squash(pdftext[i]),('PDF title missing',i+1)
        for shape in slide.shapes:
            assert shape.left>=0 and shape.top>=0 and shape.left+shape.width<=presentation.slide_width+20 and shape.top+shape.height<=presentation.slide_height+20,(i+1,shape.name)
    groups=list(csv.DictReader((ROOT/'reports/oct03_round7/group_metrics.csv').open(encoding='utf-8')))
    allrows={r['policy']:r for r in groups if r['dimension']=='all'}
    policies=['FIFO','Q','GQ','RANDOM','PROCESS_DEVIATION']
    capture=[float(allrows[p]['hits_mean'])/int(allrows[p]['positive'])*100 for p in policies]
    charts=[shape.chart for slide in presentation.slides for shape in slide.shapes if shape.has_chart]
    assert len(charts)==2
    assert all(math.isclose(a,b,abs_tol=1e-10) for a,b in zip(charts[0].series[0].values,capture)),'Fraction/percent chart mismatch'
    assert charts[0].value_axis.minimum_scale==0 and charts[0].value_axis.maximum_scale==25
    for value in capture:assert f'{value:.3f}' in pdftext[4],('PDF capture',value)
    changes=[float(next(r for r in groups if r['dimension']=='run_id' and r['group']==str(i) and r['policy']=='Q')['delta_hits_vs_FIFO']) for i in range(1,5)]
    assert all(math.isclose(a,b,abs_tol=1e-10) for a,b in zip(charts[1].series[0].values,changes))
    assert math.isclose(sum(changes),-7.6,abs_tol=1e-10)
    report=json.loads((ROOT/'reports/submission_round8/report_evidence.json').read_text(encoding='utf-8'))['metrics']
    for key in ['Q_CAPTURE','GQ_CAPTURE']:
        expected=report[key]['display']+'%';assert expected in texts[5] and expected in pdftext[5],(key,expected)
    for key in ['GQ_FPR','OOD05_FPR','GQ_INTERVENTION','OOD05_INTERVENTION','GQ_CAPTURE','OOD05_CAPTURE']:
        expected=report[key]['display']+'%';assert expected in texts[6] and expected in pdftext[6],(key,expected)
    q1=list(csv.DictReader((ROOT/'reports/oct03_pilot_r2/q1_comparison.csv').open(encoding='utf-8')))
    for fold in ['2','6']:
        for budget in [.1,.2,.3]:
            subset=[r for r in q1 if r['scheme']=='forward_run' and r['fold']==fold and float(r['budget'])==budget]
            assert len(subset)==5
            for key in ['capture_base','capture']:
                value=sum(float(r[key]) for r in subset)/5*100
                assert f'{value:.3f}%' in texts[13] and f'{value:.3f}%' in pdftext[13]
    for item in meta['numeric_metrics'].values():assert digest(ROOT/item['source'])==item['sha256']
    with zipfile.ZipFile(pptx) as z:
        for name in z.namelist():
            if name.endswith('.xml'):
                content=z.read(name).decode('utf-8');assert not re.search(r'Users[\\/]+JH|C:\\Users|C:/Users',content,re.I),(name,'personal path')
        assert sum(x.startswith('ppt/notesSlides/notesSlide') and x.endswith('.xml') for x in z.namelist())==17
    assert not presentation.core_properties.author and not presentation.core_properties.last_modified_by
    assert not (pages.metadata or {}).get('/Author')
    native=json.loads((OUT/'checks/native_export.json').read_text(encoding='utf-8-sig'))
    assert native['renderer']=='Microsoft PowerPoint' and native['slides']==17 and native['overflow_count']==0
    assert len(list((OUT/'rendered').glob('slide-*.png')))==17
    for sid,paths in meta['source_index'].items():
        for path in paths.split('; '):assert (ROOT/path).exists(),(sid,path)
    result={'resolved_historical_sources':resolved_sources,'checked_at':datetime.now(timezone.utc).isoformat(),'slides':17,'speaker_notes':17,'native_charts':2,
      'source_files':len(meta['source_hashes']),'bound_metrics':len(meta['numeric_metrics']),
      'capture_chart_recomputed_from_numerators_denominators':True,'percent_units_and_pdf_text_verified':True,
      'q1_budget_table_recomputed':True,'notes_match_script':True,'native_powerpoint_overflow':0,
      'author_blank_and_known_personal_paths_absent':True,'full_blind_submission_clearance':False,
      'new_training':False,'new_performance_evidence':False,
      'files':{p.name:{'bytes':p.stat().st_size,'sha256':digest(p)} for p in [pptx,pdf,OUT/'SPEAKER_NOTES.md',OUT/'JUDGE_QA.md']}}
    return result
if __name__=='__main__':
    result=verify();(OUT/'checks/artifact_verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps(result,ensure_ascii=False))
