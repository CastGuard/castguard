"""Verify current editable/report exports and consequential wording; no PDF regeneration."""
from pathlib import Path
from hashlib import sha256
import argparse,csv,json,math,re,zipfile,sys
from pptx import Presentation
from pypdf import PdfReader

ROOT=Path(__file__).resolve().parent
if str(ROOT) not in sys.path:sys.path.insert(0,str(ROOT))
def digest(p):return sha256(p.read_bytes()).hexdigest()
def compact(s):return re.sub(r'\s+','',s)


def verify(root=ROOT,review_folder='review'):
    root=Path(root).resolve();out=(root/review_folder).resolve()
    if not out.is_relative_to(root) or out==root:raise ValueError('Review folder must be inside this project')
    report=json.loads((out/'report_evidence.json').read_text(encoding='utf-8'))
    render=json.loads((out/'render_receipt.json').read_text(encoding='utf-8'))
    text=(out/'CastGuard_report.md').read_text(encoding='utf-8')
    from review_claims import validate
    claim_checks=validate(root,text)
    pdf=PdfReader(out/'CastGuard_report.pdf');pdftext='\n'.join(p.extract_text() for p in pdf.pages)
    assert digest(out/'CastGuard_report.md')==report['report_sha256']==render['markdown_sha256']
    assert digest(out/'CastGuard_report.pdf')==render['pdf_sha256']
    assert digest(out/'CastGuard_report.html')==render['html_sha256']
    for name,item in report['metrics'].items():
        assert digest(root/item['source'])==item['sha256'],name
        assert item['display'] in text,name
    for claim in ['불량 위험도가 높은 Shot','유효 기준선 하나','범위는0을 포함','품질검사 평균 대기는 FIFO 310.308기록, Q 142.253기록','revision 3','review.py ready']:
        assert compact(claim) in compact(text) and compact(claim) in compact(pdftext),claim
    assert '품질 가능성이 높은' not in text and '성능 작업을 중단했다' not in text
    presentation=Presentation(out/'CastGuard_presentation.pptx')
    slides=PdfReader(out/'CastGuard_presentation.pdf')
    meta=json.loads((out/'slide_evidence.json').read_text(encoding='utf-8'))
    assert len(presentation.slides)==len(slides.pages)==17
    assert digest(out/'CastGuard_presentation.pptx')==meta['pptx_sha256']
    for name,h in meta['source_hashes'].items():assert digest(root/name)==h,name
    for item in meta['numeric_metrics'].values():assert digest(root/item['source'])==item['sha256']
    contents=[]
    for i,(slide,page) in enumerate(zip(presentation.slides,slides.pages)):
        ppttext='\n'.join(s.text for s in slide.shapes if s.has_text_frame)
        assert meta['slides'][i]['title'] in ppttext
        assert compact(meta['slides'][i]['title']) in compact(page.extract_text())
        assert slide.notes_slide.notes_text_frame.text==meta['slides'][i]['speaker']
        contents.append(ppttext)
    note=meta['slides'][4]['speaker'];assert all(x in note for x in ['유효 기준선 하나','범위는0을 포함','대기가 더 깁니다','통계적 우위'])
    assert '표본 선언 검사' in contents[9] and '계획 파일·해시 일치' in contents[9]
    assert '실제 독립성은 원천 검토' in contents[9]
    assert 'revision 3' in meta['slides'][9]['speaker'] and '인증하지 않습니다' in meta['slides'][9]['speaker']
    groups=list(csv.DictReader((root/'reports/oct03_round7/group_metrics.csv').open(encoding='utf-8')))
    allrows={r['policy']:r for r in groups if r['dimension']=='all'}
    charts=[s.chart for slide in presentation.slides for s in slide.shapes if s.has_chart]
    assert len(charts)==2
    policies=['FIFO','Q','GQ','RANDOM','PROCESS_DEVIATION']
    expected=[100*float(allrows[p]['hits_mean'])/int(allrows[p]['positive']) for p in policies]
    for actual,value in zip(charts[0].series[0].values,expected):assert math.isclose(actual,value,abs_tol=1e-10)
    assert charts[0].value_axis.minimum_scale==0 and charts[0].value_axis.maximum_scale==25
    for value in expected:assert f'{value:.3f}' in slides.pages[4].extract_text()
    qa=(out/'JUDGE_QA.md').read_text(encoding='utf-8');notes=(out/'SPEAKER_NOTES.md').read_text(encoding='utf-8')
    assert '109개' not in qa+notes and '공식 발표시간이 아니다' in notes
    assert not presentation.core_properties.author and not presentation.core_properties.last_modified_by
    assert not (pdf.metadata or {}).get('/Author') and not (slides.metadata or {}).get('/Author')
    with zipfile.ZipFile(out/'CastGuard_presentation.pptx') as archive:
        for name in archive.namelist():
            if name.endswith('.xml'):assert not re.search(r'Users[\\/]+JH|C:\\Users|C:/Users',archive.read(name).decode('utf-8'),re.I),name
    native=json.loads((out/'checks/native_export.json').read_text(encoding='utf-8-sig'))
    assert native['renderer']=='Microsoft PowerPoint' and native['slides']==17 and native['overflow_count']==0
    return {'report_pages':len(pdf.pages),'report_metric_bindings':len(report['metrics']),
        'focused_claim_checks':claim_checks,
        'presentation_slides':17,'slide_metric_bindings':len(meta['numeric_metrics']),'native_powerpoint_overflow':0,
        'one_effective_baseline_zero_inclusive_sensitivity_and_delay_caveats_verified':True,
        'revision3_scope_and_readiness_entrypoint_verified':True,
        'numeric_results_unchanged':True,'external_prerequisites_remain':True,'new_performance_evidence':False}


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--review-folder',default='review');args=parser.parse_args()
    print(json.dumps(verify(review_folder=args.review_folder),ensure_ascii=False,indent=2))
