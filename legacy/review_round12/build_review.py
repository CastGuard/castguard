"""Build the current report/presentation from preserved evidence; never overwrite history."""
from pathlib import Path
from hashlib import sha256
from datetime import datetime,timezone
import argparse,json,importlib.util,sys
from pptx import Presentation

ROOT=Path(__file__).resolve().parent
if str(ROOT) not in sys.path:sys.path.insert(0,str(ROOT))
def digest(path):return sha256(path.read_bytes()).hexdigest()


def build(output):
    output=Path(output).resolve()
    if not output.is_relative_to(ROOT) or output==ROOT:raise ValueError('Use a new folder inside this project')
    missing_fonts=[name for name in ['batang.ttc','malgunbd.ttf'] if not (Path('C:/Windows/Fonts')/name).is_file()]
    if missing_fonts:raise RuntimeError('Report re-render needs installed Windows fonts: '+', '.join(missing_fonts)+'. The included PDF can be viewed without re-rendering.')
    output.mkdir(parents=True,exist_ok=False);prefix=output.relative_to(ROOT).as_posix()
    (output/'checks').mkdir();(output/'evidence').mkdir()
    lineage_name='reports/oct03_evidence/final/lineage_audit.json'
    lineage=json.loads((ROOT/lineage_name).read_text(encoding='utf-8'));lineage['source']='data/raw'
    local_lineage=output/'evidence/lineage_audit.json'
    local_lineage.write_text(json.dumps(lineage,ensure_ascii=False,indent=2),encoding='utf-8')
    def bind_export_lineage(items):
        for item in items.values():
            if item['source']==lineage_name:
                item.update(source=prefix+'/evidence/lineage_audit.json',sha256=digest(local_lineage))
    prior=ROOT/'reports/submission_round8/CastGuard_current_review.md'
    text=prior.read_text(encoding='utf-8')
    text=text.replace('회차8 최신 검토용 편집초안:','최신 통합 검토용 편집초안:')
    text=text.replace('품질 가능성이 높은 Shot','불량 위험도가 높은 Shot')
    text=text.replace('현재는 새 데이터와 사용자 운영조건이 필요한 지점에서 성능 작업을 중단했다.',
        '새 성능 선택은 독립 자료와 사전 설계가 확보될 때까지 보류한다. 코드·재현·심사 설명은 계속 보완한다.')
    text=text.replace('미래 평가용 JSON 계약과 빈 CSV8개를 만들었다.',
        '미래 평가용 JSON 계약과 빈 CSV8개를 만들었고 현재 기본 검증기는 revision 3이다. future_evaluation.py와 기존 v2 명령 모두 같은 구현을 실행한다.')
    text=text.replace('원천 해시와 원본 event_id가 과거자료와 같으면 새 평가로 거부한다.',
        '원천 해시와 event_id/원본 event_id가 과거자료와 같으면 새 평가로 거부한다. 점수 결정은 검사 시작보다 엄격히 앞서야 한다. 실제자료는 고정 추출계획 파일·해시·계약 값 일치까지 검사하지만 실제 감사 독립성은 사람의 원천 검토가 필요하다.')
    text=text.replace('python replay_report.py --full-evidence\npython run_review_tests.py\npython build_report.py --check\npython verify_judge_baselines.py\npython verify_round7.py',
        'python -I review.py check --full-replay --require-package\npython -I review.py metrics\npython -I review.py ready')
    text=text.replace('회차6 단순 기준은 별도 검증 명령과 근거 파일을 사용한다.',
        '최신 검토 명령은 원본 불량열과 행별 행동에서 주요 지표를 독립 계산하고, 전체 회귀시험과 중요한 오류8종을 주입한 변이시험을 수행한다. 과거빈도와 FIFO는 같은 유효 기준선 하나로 다룬다.')
    text=text.replace('`ROUND8_README.md`','`CURRENT_REVIEW.md`')
    text=text.replace('실제 팀 정보·서명·설문 완료 화면, 최종 양식 적용, 발표 PDF/PPT, 블라인드 검사 및 포털 접수 증빙은 별도로 완료해야 한다.',
        '실제 팀 정보·서명·설문 완료 화면, 최종 양식 적용, 블라인드 검사 및 포털 접수 증빙은 별도로 완료해야 한다. 검토용 발표 PDF/PPTX는 함께 제공하며 PPTX 허용 여부는 포털에서 확인해야 한다. review.py ready가 제출 준비와 미래 성능평가의 부족한 증거를 따로 출력한다. 기능 검사 통과가 접수 완료를 뜻하지 않는다.')
    text=text.replace('이 자료에서 Q20.871%와 GQ20.238%보다 높았다는 사실은 FIFO의 현장 우위나 미래 일반화를 입증하지 않는다.',
        '관측 점추정은 Q20.871%와 GQ20.238%보다 높지만, 아래 구간 재표본 범위는0을 포함하고 대기 부담도 다르다. 통계적으로 확립된 FIFO 우위·모델 열등성 또는 어느 정책의 현장 우위도 입증하지 않는다.')
    (output/'CastGuard_current_review.md').write_text(text,encoding='utf-8')
    # Reuse the proven renderer layout, changing only the current label and cover
    # prerequisite wording. HWPX is an optional editing format, not a required PDF.
    source=(ROOT/'render_submission_round8.py').read_text(encoding='utf-8')
    source=source.replace('회차8 최신 공식 양식 대응 검토초안','현재 통합 공식 양식 대응 검토초안')
    source=source.replace('실제 서명·설문·최종 HWPX 변환·제출 미완료','실제 팀정보·서명·설문·최종 제출 미완료')
    namespace={'__file__':str(ROOT/'render_submission_round8.py'),'__name__':'review_report_renderer'}
    exec(compile(source,'render_submission_round8.py','exec'),namespace)
    namespace['OUT']=output;namespace['render']()
    for ext in ['md','html','pdf']:(output/('CastGuard_current_review.'+ext)).rename(output/('CastGuard_report.'+ext))
    index=json.loads((ROOT/'reports/submission_round8/report_evidence.json').read_text(encoding='utf-8'))
    bind_export_lineage(index['metrics'])
    report_index={'created_at':datetime.now(timezone.utc).isoformat(),'historical_report_source':prior.relative_to(ROOT).as_posix(),
        'historical_report_sha256':digest(prior),'report_sha256':digest(output/'CastGuard_report.md'),
        'metrics':index['metrics'],'numeric_results_unchanged':True,'new_performance_evidence':False,
        'changes':['clear defect-risk wording','one effective FIFO/history baseline','zero-inclusive sensitivity and delay caveats',
                   'revision 3 default and honest source-file scope','one executable review/readiness entrypoint']}
    (output/'report_evidence.json').write_text(json.dumps(report_index,ensure_ascii=False,indent=2),encoding='utf-8')

    original=ROOT/'reports/presentation_round9';presentation=Presentation(original/'CastGuard_presentation_draft.pptx')
    meta=json.loads((original/'slide_evidence.json').read_text(encoding='utf-8'))
    bind_export_lineage(meta['numeric_metrics'])
    replacements={
        'S11 future evaluation contract / S12 v2 regression tests':'S11 input contract / S12 revision 3 & mutation tests',
        '② 독립 표본':'② 표본 선언 검사',
        '전체 생산행 유지\n고정확률 추출 확인\n모든 감사 표본의 정답\n정책 선택 표본 혼합 금지':'전체 생산행 유지\n고정 추출 재계산\n계획 파일·해시 일치\n실제 독립성은 원천 검토',
    }
    for slide in presentation.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:continue
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    for old,new in replacements.items():run.text=run.text.replace(old,new)
    # Multi-paragraph card text is stored as separate runs. Change text per line,
    # retaining the existing fonts, size and native shapes.
    line_replacements={'고정된 감사 추출':'고정 추출 재계산','모든 선택 표본의 정답':'계획 파일·해시 일치','정책 선택 표본 혼입 차단':'실제 독립성은 원천 검토'}
    for shape in presentation.slides[9].shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    for old,new in line_replacements.items():run.text=run.text.replace(old,new)
    meta['slides'][4]['speaker']='관측된 포착률 점추정은 FIFO22.376%, Q20.871%, GQ20.238%입니다. 과거 제품빈도는 구간 안에서 상수 점수여서 FIFO와 검사 행과 시점이 완전히 같았습니다. 따라서 유효 기준선 하나로 다룹니다. 뒤에서 보듯 네 구간 재표본 범위는0을 포함하고 FIFO는 품질검사 대기가 더 깁니다. 이 값만으로 통계적 우위, 모델의 일반적인 열등성 또는 현장 우위를 판단하지 않습니다.'
    meta['slides'][9]['speaker']='기본 명령과 기존 v2 명령은 모두 revision 3을 실행합니다. 미래 입력과 평가 라벨 사용, 검사 시작 동시각, 검사 완료와 만료·대기한도 경계를 검사합니다. 실제자료에는 추출계획 파일의 해시와 계약 일치를 요구합니다. 이 검사는 선언된 시각과 로컬 파일의 일치이며 실제 감사 독립성이나 시계 진위를 인증하지 않습니다. 정상 회귀시험과 의도적으로 검사를 약화한8개 변이시험을 구분해 확인했고, 기능시험을 모델 성능 근거로 쓰지 않습니다.'
    meta['slides'][9]['source_ids']='S11 input contract / S12 revision 3 & mutation tests'
    for item,slide in zip(meta['slides'],presentation.slides):slide.notes_slide.notes_text_frame.text=item['speaker']
    presentation.core_properties.author='';presentation.core_properties.last_modified_by=''
    presentation.save(output/'CastGuard_presentation.pptx')
    # Store a path-normalized evidence copy for current citations. The historic
    # project lineage JSON stays unchanged; numerical fields are identical.
    lineage_path=ROOT/'reports/oct03_evidence/final/lineage_audit.json'
    lineage=json.loads(lineage_path.read_text(encoding='utf-8'));lineage['source']='data/raw'
    local_lineage=output/'evidence/lineage_audit.json';local_lineage.write_text(json.dumps(lineage,ensure_ascii=False,indent=2),encoding='utf-8')
    meta['source_index']['S02']=prefix+'/CastGuard_report.md'
    meta['source_index']['S04']=prefix+'/evidence/lineage_audit.json'
    meta['source_index']['S12']='future_evaluation_core.py; tests/test_future_evaluation_adversarial.py; review_mutations.py'
    meta['source_index']['S16']='docs/COMPETITION_READINESS_ROUND10.md'
    sources={}
    for paths in meta['source_index'].values():
        for path in paths.split('; '):
            if (ROOT/path).is_file():sources[path]=digest(ROOT/path)
    meta.update(created_at=datetime.now(timezone.utc).isoformat(),source_hashes=sources,
        pptx_sha256=digest(output/'CastGuard_presentation.pptx'),validator_revision=3,
        preserved_source_presentation_sha256=digest(original/'CastGuard_presentation_draft.pptx'))
    (output/'slide_evidence.json').write_text(json.dumps(meta,ensure_ascii=False,indent=2),encoding='utf-8')
    notes=['# CastGuard 발표자 원고 — 현재 통합 검토본','',
        '본문13장+부록4장. 본문600초는 연습 배분이며 공식 발표시간이 아니다. 같은 원고가 PPTX 노트에 있다.','']
    for item in meta['slides']:
        notes += [f'## {item["number"]:02d}. {item["title"]}',f'연습 배분: {item["practice_seconds"]}초' if not item['appendix'] else '질문 시 쓰는 부록',item['speaker'],'']
    (output/'SPEAKER_NOTES.md').write_text('\n'.join(notes),encoding='utf-8')
    qa=(original/'JUDGE_QA.md').read_text(encoding='utf-8').replace('109개 테스트는 합성 기능 검증을 포함한 소프트웨어 확인이지 성능 실적이 아닙니다.',
        '현재 기본 검증기는 revision 3이며 추출계획 파일의 해시·값 일치도 요구합니다. 전체 회귀시험과 의도적 오류8종 변이시험은 소프트웨어 확인이며 성능 실적이 아닙니다.')
    qa=qa.replace('COMPETITION_REVIEW.md','../docs/COMPETITION_READINESS_ROUND10.md').replace('../../docs/','../docs/')
    (output/'JUDGE_QA.md').write_text(qa,encoding='utf-8')
    from review_readiness import template
    (output/'readiness_evidence.template.json').write_text(json.dumps(template(),ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'output':prefix,'report_metric_bindings':len(index['metrics']),'slides':len(presentation.slides),'validator_revision':3}))


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--output',type=Path,default=ROOT/'review');args=parser.parse_args();build(args.output)
