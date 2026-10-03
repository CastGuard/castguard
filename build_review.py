"""Build the current report/presentation from preserved evidence; never overwrite history."""
from pathlib import Path
from hashlib import sha256
from datetime import datetime,timezone
import argparse,json,importlib.util,sys
from pptx import Presentation

ROOT=Path(__file__).resolve().parent
if str(ROOT) not in sys.path:sys.path.insert(0,str(ROOT))
def digest(path):return sha256(path.read_bytes()).hexdigest()


def build(output):
    output=Path(output).resolve()
    if not output.is_relative_to(ROOT) or output==ROOT:raise ValueError('Use a new folder inside this project')
    missing_fonts=[name for name in ['batang.ttc','malgunbd.ttf'] if not (Path('C:/Windows/Fonts')/name).is_file()]
    if missing_fonts:raise RuntimeError('Report re-render needs installed Windows fonts: '+', '.join(missing_fonts)+'. The included PDF can be viewed without re-rendering.')
    output.mkdir(parents=True,exist_ok=False);prefix=output.relative_to(ROOT).as_posix()
    (output/'checks').mkdir();(output/'evidence').mkdir()
    lineage_name='reports/oct03_evidence/final/lineage_audit.json'
    lineage=json.loads((ROOT/lineage_name).read_text(encoding='utf-8'));lineage['source']='data/raw'
    local_lineage=output/'evidence/lineage_audit.json'
    local_lineage.write_text(json.dumps(lineage,ensure_ascii=False,indent=2),encoding='utf-8')
    def bind_export_lineage(items):
        for item in items.values():
            if item['source']==lineage_name:
                item.update(source=prefix+'/evidence/lineage_audit.json',sha256=digest(local_lineage))
    prior=ROOT/'reports/submission_round8/CastGuard_current_review.md'
    text=prior.read_text(encoding='utf-8')
    text=text.replace('회차8 최신 검토용 편집초안:','최신 통합 검토용 편집초안:')
    text=text.replace('품질 가능성이 높은 Shot','불량 위험도가 높은 Shot')
    text=text.replace('양식 대응 검토초안: 실제 서명·설문·최종 HWPX 변환·제출 미완료.',
        '양식 대응 검토초안: 실제 팀정보·서명·설문·최종 제출 미완료.')
    text=text.replace('표의 모집단은 1,387 Shot, 양성 272개이며 값은 seed별 지표의 평균이다. B-A AP 변화 0.00535는 사전 개선 기준 0.02에 못 미친다. A0가 A보다 높았다는 점도 함께 공개한다. 센서나 이력을 더하면 성능이 반드시 좋아지는 것은 아니다. 평균 확률을 먼저 낸 뒤 재표집한 AP 차이와 이 표의 seed 지표 평균 차이는 서로 다른 추정량이므로 혼용하지 않는다.[E01]',
        '이 표는 과거 test 1,387 Shot(양성 272개)의 seed 지표 평균이다. B-A AP 차이 0.00535는 기술값이며 선택 판정에 쓰지 않는다. 별도 validation의 B-A는 -0.00454로 사전 기준 0.02에 미달했다. A0가 A보다 높았다는 점도 공개한다. 평균 확률을 재표집한 AP 차이와 seed 지표 평균 차이는 서로 다른 추정량이다.[E01]')
    text=text.replace('E01 `reports/oct02/ablation_summary.csv`, `bootstrap_intervals.csv`:',
        'E01 `reports/oct02/ablation_summary.csv`, `bootstrap_intervals.csv`, `history_selection.json`:')
    text=text.replace('현재는 새 데이터와 사용자 운영조건이 필요한 지점에서 성능 작업을 중단했다.',
        '새 성능 선택은 독립 자료와 사전 설계가 확보될 때까지 보류한다. 코드·재현·심사 설명은 계속 보완한다.')
    text=text.replace('미래 평가용 JSON 계약과 빈 CSV8개를 만들었다.',
        '미래 평가용 JSON 계약과 빈 CSV8개를 만들었고 현재 기본 검증기는 revision 3이다. future_evaluation.py와 기존 v2 명령 모두 같은 구현을 실행한다.')
    text=text.replace('원천 해시와 원본 event_id가 과거자료와 같으면 새 평가로 거부한다.',
        '원천 해시와 event_id/원본 event_id가 과거자료와 같으면 새 평가로 거부한다. 점수 결정은 검사 시작보다 엄격히 앞서야 한다. 실제자료는 고정 추출계획 파일·해시·계약 값 일치까지 검사하지만 실제 감사 독립성은 사람의 원천 검토가 필요하다.')
    text=text.replace('python replay_report.py --full-evidence\npython run_review_tests.py\npython build_report.py --check\npython verify_judge_baselines.py\npython verify_round7.py',
        'python -I review.py check --full-replay --require-package\npython -I review.py metrics\npython -I review.py ready')
    text=text.replace('회차6 단순 기준은 별도 검증 명령과 근거 파일을 사용한다.',
        '최신 검토 명령은 원본 불량열과 행별 행동에서 주요 지표를 독립 계산하고, 전체 회귀시험과 중요한 오류8종을 주입한 변이시험을 수행한다. 과거빈도와 FIFO는 같은 유효 기준선 하나로 다룬다.')
    text=text.replace('`ROUND8_README.md`','`CURRENT_REVIEW.md`')
    text=text.replace('실제 팀 정보·서명·설문 완료 화면, 최종 양식 적용, 발표 PDF/PPT, 블라인드 검사 및 포털 접수 증빙은 별도로 완료해야 한다.',
        '실제 팀 정보·서명·설문 완료 화면, 최종 양식 적용, 블라인드 검사 및 포털 접수 증빙은 별도로 완료해야 한다. 검토용 발표 PDF/PPTX는 함께 제공하며 PPTX 허용 여부는 포털에서 확인해야 한다. review.py ready가 제출 준비와 미래 성능평가의 부족한 증거를 따로 출력한다. 기능 검사 통과가 접수 완료를 뜻하지 않는다.')
    text=text.replace('이 자료에서 Q20.871%와 GQ20.238%보다 높았다는 사실은 FIFO의 현장 우위나 미래 일반화를 입증하지 않는다.',
        '관측 점추정은 Q20.871%와 GQ20.238%보다 높지만, 아래 구간 재표본 범위는0을 포함하고 대기 부담도 다르다. 통계적으로 확립된 FIFO 우위·모델 열등성 또는 어느 정책의 현장 우위도 입증하지 않는다.')
    (output/'CastGuard_current_review.md').write_text(text,encoding='utf-8')
    # Reuse the proven renderer layout, changing only the current label and cover
    # prerequisite wording. HWPX is an optional editing format, not a required PDF.
    source=(ROOT/'render_submission_round8.py').read_text(encoding='utf-8')
    source=source.replace('회차8 최신 공식 양식 대응 검토초안','현재 통합 공식 양식 대응 검토초안')
    source=source.replace('실제 서명·설문·최종 HWPX 변환·제출 미완료','실제 팀정보·서명·설문·최종 제출 미완료')
    namespace={'__file__':str(ROOT/'render_submission_round8.py'),'__name__':'review_report_renderer'}
    exec(compile(source,'render_submission_round8.py','exec'),namespace)
    namespace['OUT']=output;namespace['render']()
    for ext in ['md','html','pdf']:(output/('CastGuard_current_review.'+ext)).rename(output/('CastGuard_report.'+ext))
    index=json.loads((ROOT/'reports/submission_round8/report_evidence.json').read_text(encoding='utf-8'))
    selection_name='reports/oct02/history_selection.json'
    selected=next(x for x in json.loads((ROOT/selection_name).read_text(encoding='utf-8')) if x['scheme']=='within_run' and x['fold']=='primary' and x['model']=='lightgbm')
    validation_metric={'value':selected['b_validation_ap_gain'],'display':f"{selected['b_validation_ap_gain']:.5f}",
        'unit':'validation AP difference; mean of five paired seeds','source':selection_name,
        'selector':'scheme=within_run;fold=primary;model=lightgbm;b_validation_ap_gain','sha256':digest(ROOT/selection_name)}
    index['metrics']['VALIDATION_B_AP_GAIN']=validation_metric
    bind_export_lineage(index['metrics'])
    report_index={'created_at':datetime.now(timezone.utc).isoformat(),'historical_report_source':prior.relative_to(ROOT).as_posix(),
        'historical_report_sha256':digest(prior),'report_sha256':digest(output/'CastGuard_report.md'),
        'metrics':index['metrics'],'numeric_results_unchanged':True,'new_performance_evidence':False,
        'changes':['clear defect-risk wording','one effective FIFO/history baseline','zero-inclusive sensitivity and delay caveats',
                   'revision 3 default and honest source-file scope','one executable review/readiness entrypoint']}
    (output/'report_evidence.json').write_text(json.dumps(report_index,ensure_ascii=False,indent=2),encoding='utf-8')

    original=ROOT/'reports/presentation_round9';presentation=Presentation(original/'CastGuard_presentation_draft.pptx')
    meta=json.loads((original/'slide_evidence.json').read_text(encoding='utf-8'))
    bind_export_lineage(meta['numeric_metrics'])
    replacements={
        'S02 §2·§3 / S08 ablation_summary / S05 group_metrics':'S02 §2 / S08 test AP / S17 validation / S05 queue',
        'within_run AP':'과거 test AP',
        '증가 기준 +0.02 미달':'차이 +0.00535 (기술값)',
        '통합 구현과 통합 성능 개선을 구분했습니다. 원본·분할·실패 결과를 보존합니다.':
        '선택은 별도 validation: B-A AP -0.00454, 기준 +0.02 미달',
        '개발 선택은 안쪽 시간분할에서 수행했습니다. 기존 자료는 이미 노출되어 탐색 평가입니다.':
        'FIFO: 선입순 · Q: 품질 위험도 · GQ: 품질+상태 | 이미 본 자료의 탐색 평가',
        'S11 future evaluation contract / S12 v2 regression tests':'S11 input contract / S12 revision 3 & mutation tests',
        '② 독립 표본':'② 표본 선언 검사',
        '전체 생산행 유지\n고정확률 추출 확인\n모든 감사 표본의 정답\n정책 선택 표본 혼합 금지':'전체 생산행 유지\n고정 추출 재계산\n계획 파일·해시 일치\n실제 독립성은 원천 검토',
    }
    for slide in presentation.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:continue
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    for old,new in replacements.items():run.text=run.text.replace(old,new)
    # Multi-paragraph card text is stored as separate runs. Change text per line,
    # retaining the existing fonts, size and native shapes.
    line_replacements={'고정된 감사 추출':'고정 추출 재계산','모든 선택 표본의 정답':'계획 파일·해시 일치','정책 선택 표본 혼입 차단':'실제 독립성은 원천 검토'}
    for shape in presentation.slides[9].shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    for old,new in line_replacements.items():run.text=run.text.replace(old,new)
    meta['slides'][4]['speaker']='관측된 포착률 점추정은 FIFO22.376%, Q20.871%, GQ20.238%입니다. 과거 제품빈도는 구간 안에서 상수 점수여서 FIFO와 검사 행과 시점이 완전히 같았습니다. 따라서 유효 기준선 하나로 다룹니다. 뒤에서 보듯 네 구간 재표본 범위는0을 포함하고 FIFO는 품질검사 대기가 더 깁니다. 이 값만으로 통계적 우위, 모델의 일반적인 열등성 또는 현장 우위를 판단하지 않습니다.'
    meta['slides'][9]['speaker']='기본 명령과 기존 v2 명령은 모두 revision 3을 실행합니다. 미래 입력과 평가 라벨 사용, 검사 시작 동시각, 검사 완료와 만료·대기한도 경계를 검사합니다. 실제자료에는 추출계획 파일의 해시와 계약 일치를 요구합니다. 이 검사는 선언된 시각과 로컬 파일의 일치이며 실제 감사 독립성이나 시계 진위를 인증하지 않습니다. 정상 회귀시험과 의도적으로 검사를 약화한8개 변이시험을 구분해 확인했고, 기능시험을 모델 성능 근거로 쓰지 않습니다.'
    meta['slides'][9]['source_ids']='S11 input contract / S12 revision 3 & mutation tests'
    # Keep one canonical deck; clarify the spoken definitions and scope without new numeric claims.
    meta['slides'][3]['speaker']='FIFO는 먼저 들어온 순서, Q는 품질 위험도 순서, GQ는 품질과 설비 상태를 결합한 정책입니다. 안쪽 시간순 분할에서 후보를 고정하고, 이미 본 네 바깥 구간에서 비교했습니다. 이 표의 비교 범위는 원본 전체가 아니라 생산3395행입니다. 같은 누적 상한에서도 실제 검사는 Q와FIFO675건, GQ677건이었습니다. 포착률의 분모는 품질 관측2879행 중 양성505개입니다. 정답 없는516행도 검사 부담에서 빼지 않았습니다.'
    meta['slides'][5]['speaker']='첫 행은 불량 위험도 순위의 평균정밀도, AP입니다. 과거 테스트에서 A는.2614, 이력 결합B는.2668이며 차이는 기술값입니다. 선택 판정은 별도 validation에서 했고 B-A는 마이너스.00454로 기준.02에 미달했습니다. 두 번째 순차 검사에서는 Q에서GQ로 결합할 때 관측 포착률이.634퍼센트포인트 낮았습니다. 서로 다른 평가를 합산하거나 test 수치로 모델을 고르지 않습니다.'
    meta['slides'][9]['speaker']='화면의 세 가지를 검사합니다. 첫째, 점수를 만들 때 아직 도착하지 않은 입력이나 평가 정답을 사용했는지 봅니다. 둘째, 실제자료의 고정 추출계획 파일과 해시·계약 값이 일치하는지 확인합니다. 셋째, 검사 자원과 대기·만료 경계를 지켰는지 봅니다. 현재 기본 검증기는 revision 3입니다. 이것은 선언된 시각과 파일의 일치 검사이며 실제 감사 독립성이나 시계 진위를 인증하지 않습니다. 기능시험을 모델 성능 실적으로 쓰지 않습니다.'
    meta['slides'][10]['speaker']='현재 기여는 제조 데이터 결합을 검사 자원과 대기를 명시한 조건에서 검증하는 실험 체계입니다. 품질과 상태의 결합 전후, 상태 경보가 품질 검사 자원을 쓰는 기회비용, 정답 없는 행의 부담을 함께 계산했습니다. 실제 기업문제 해결이나 경제효과는 아직 입증하지 못했습니다. 이를 다음 실증에서 채워야 할 핵심 공백으로 남겼습니다.'
    # Provisional 10-minute allocation only; official speaking/Q&A limits remain unknown.
    rehearsal_seconds=[35,40,45,60,55,50,40,50,45,55,40,50,35]
    assert sum(rehearsal_seconds)==600
    for item,seconds in zip(meta['slides'][:13],rehearsal_seconds):item['practice_seconds']=seconds
    for item,slide in zip(meta['slides'],presentation.slides):slide.notes_slide.notes_text_frame.text=item['speaker']
    presentation.core_properties.author='';presentation.core_properties.last_modified_by=''
    presentation.save(output/'CastGuard_presentation.pptx')
    # Store a path-normalized evidence copy for current citations. The historic
    # project lineage JSON stays unchanged; numerical fields are identical.
    lineage_path=ROOT/'reports/oct03_evidence/final/lineage_audit.json'
    lineage=json.loads(lineage_path.read_text(encoding='utf-8'));lineage['source']='data/raw'
    local_lineage=output/'evidence/lineage_audit.json';local_lineage.write_text(json.dumps(lineage,ensure_ascii=False,indent=2),encoding='utf-8')
    meta['source_index']['S02']=prefix+'/CastGuard_report.md'
    meta['source_index']['S04']=prefix+'/evidence/lineage_audit.json'
    meta['source_index']['S12']='future_evaluation_core.py; tests/test_future_evaluation_adversarial.py; review_mutations.py'
    meta['source_index']['S16']='docs/COMPETITION_READINESS_ROUND10.md'
    meta['source_index']['S17']=selection_name
    meta['numeric_metrics']['VALIDATION_B_AP_GAIN']={**validation_metric,'slide_multiplier':1}
    meta['slides'][5]['source_ids']+=' / S17 validation selection'
    sources={}
    for paths in meta['source_index'].values():
        for path in paths.split('; '):
            if (ROOT/path).is_file():sources[path]=digest(ROOT/path)
    meta.update(created_at=datetime.now(timezone.utc).isoformat(),source_hashes=sources,
        pptx_sha256=digest(output/'CastGuard_presentation.pptx'),validator_revision=3,
        preserved_source_presentation_sha256=digest(original/'CastGuard_presentation_draft.pptx'))
    (output/'slide_evidence.json').write_text(json.dumps(meta,ensure_ascii=False,indent=2),encoding='utf-8')
    notes=['# CastGuard 발표자 원고 — 현재 통합 검토본','',
        '본문13장+부록4장. 본문600초는 연습 배분이며 공식 발표시간이 아니다. 같은 원고가 PPTX 노트에 있다.','']
    for item in meta['slides']:
        notes += [f'## {item["number"]:02d}. {item["title"]}',f'연습 배분: {item["practice_seconds"]}초' if not item['appendix'] else '질문 시 쓰는 부록',item['speaker'],'']
    (output/'SPEAKER_NOTES.md').write_text('\n'.join(notes),encoding='utf-8')
    qa=(original/'JUDGE_QA.md').read_text(encoding='utf-8').replace('109개 테스트는 합성 기능 검증을 포함한 소프트웨어 확인이지 성능 실적이 아닙니다.',
        '현재 기본 검증기는 revision 3이며 추출계획 파일의 해시·값 일치도 요구합니다. 전체 회귀시험과 의도적 오류8종 변이시험은 소프트웨어 확인이며 성능 실적이 아닙니다.')
    qa=qa.replace('COMPETITION_REVIEW.md','../docs/COMPETITION_READINESS_ROUND10.md').replace('../../docs/','../docs/')
    qa=qa.replace('# 심사 예상 질문과 답변 — 회차9','# 심사 예상 질문과 답변 — 현재 검토본')
    qa=qa.replace('`../oct03_','`../reports/oct03_')
    qa=qa.replace('회차7 `fifo_history_equivalence.csv`','`../reports/oct03_round7/fifo_history_equivalence.csv`')
    qa=qa.replace('GQ는 상태 경보로 품질 미관측 행을 더 검사했고 총 검사량이2건 많았습니다.',
        '정책별 대기열과 검사 가능한 행의 차이로 총 검사량과 품질 관측 검사의 구성도 달랐습니다.')
    qa=qa.replace('인증하지 않습니다.현재','인증하지 않습니다. 현재')
    qa=qa.replace('통합 구현은 확인했지만 유효성은 목표 미달입니다.',
        '통합 구현은 확인했지만 유효성은 목표 미달입니다. 발표6장의 AP는 과거 test 기술값이며, 별도 validation의 B-A AP는 -0.00454로 개선 기준0.02에 미달했습니다. test 수치로 후보를 다시 고르지 않았습니다.')
    (output/'JUDGE_QA.md').write_text(qa,encoding='utf-8')
    from review_readiness import template
    (output/'readiness_evidence.template.json').write_text(json.dumps(template(),ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'output':prefix,'report_metric_bindings':len(index['metrics']),'slides':len(presentation.slides),'validator_revision':3}))


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--output',type=Path,default=ROOT/'review');args=parser.parse_args();build(args.output)
